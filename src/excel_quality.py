import openpyxl
from collections import Counter
from typing import List, Dict, Any, Tuple, Optional
import pandas as pd
import io

from .issue_manager import StandardizedIssue, IssueClassification

class ExcelQualityEngine:
    """
    Excel-specific quality detection (Phase 6).
    Detects inconsistent number formatting, date formatting, fill, font, and alignment.
    Generates StandardizedIssue objects that flow into the existing proposal architecture.
    """
    
    def __init__(self, raw_bytes: bytes, sheet_name: Optional[str] = None):
        self.raw_bytes = raw_bytes
        self.sheet_name = sheet_name
        self.has_formulas_or_merges = False
        
    def analyze(self, df: pd.DataFrame) -> List[StandardizedIssue]:
        issues = []
        if not self.raw_bytes:
            return []

        # Check if file is legacy binary Excel format (.xls - OLE2 compound document)
        if self.raw_bytes.startswith(b'\xd0\xcf\x11\xe0'):
            # Legacy .xls cannot be parsed by openpyxl; mark formulas/merges True conservatively to protect data
            self.has_formulas_or_merges = True
            return []

        try:
            wb = openpyxl.load_workbook(io.BytesIO(self.raw_bytes), data_only=False)
            if self.sheet_name and self.sheet_name in wb.sheetnames:
                ws = wb[self.sheet_name]
            else:
                ws = wb.active
            
            # Check for structural dependencies (formulas or merges)
            if ws.merged_cells.ranges:
                self.has_formulas_or_merges = True
            else:
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.data_type == 'f':
                            self.has_formulas_or_merges = True
                            break
                    if self.has_formulas_or_merges:
                        break
        except Exception:
            # Conservative safeguard for corrupted or unreadable Excel workbooks
            self.has_formulas_or_merges = True
            return []
            
        # Helper to get cells in a column ignoring the header (row 1)
        # We assume pandas row_index (0-based) corresponds to Excel row (2-based)
        # and pandas column_index (0-based) corresponds to Excel col (1-based)
        
        # Build mapping of column name to openpyxl column index (1-based)
        # assuming the first row is headers matching df.columns.
        header_row = [cell.value for cell in ws[1]]
        col_name_to_idx = {}
        for idx, val in enumerate(header_row):
            if val in df.columns:
                col_name_to_idx[val] = idx + 1

        # We will only process up to 1000 rows to keep it reasonably performant.
        max_rows = min(len(df), 1000)
        
        merged_coords = set()
        for mr in ws.merged_cells.ranges:
            for cell_obj in mr.cells:
                merged_coords.add(cell_obj)
        
        for col_name in df.columns:
            col_idx = col_name_to_idx.get(col_name)
            if not col_idx:
                continue
                
            # Collect formatting for the column
            formats = []
            for row_idx in range(max_rows):
                # Excel rows are 1-based, header is 1, data starts at 2
                excel_row = row_idx + 2
                cell = ws.cell(row=excel_row, column=col_idx)
                
                # Check if merged
                is_merged = (excel_row, col_idx) in merged_coords
                
                formats.append({
                    'row_index': row_idx,
                    'excel_row': excel_row,
                    'cell': cell,
                    'number_format': cell.number_format,
                    'fill': cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else None,
                    'font_name': cell.font.name if cell.font else None,
                    'font_size': cell.font.size if cell.font else None,
                    'font_bold': cell.font.bold if cell.font else None,
                    'font_italic': cell.font.italic if cell.font else None,
                    'alignment_horz': cell.alignment.horizontal if cell.alignment else None,
                    'is_merged': is_merged,
                    'is_formula': cell.data_type == 'f',
                    'is_empty': cell.value is None
                })
                
            # 1. Check Number Formats
            # We filter out empty cells from dominant format logic
            populated = [f for f in formats if not f['is_empty']]
            if populated:
                num_formats = [f['number_format'] for f in populated]
                counter = Counter(num_formats)
                if len(counter) > 1:
                    dominant_format, count = counter.most_common(1)[0]
                    # If dominant format is overwhelming (e.g. >= 90%)
                    if count / len(populated) >= 0.90:
                        for f in populated:
                            if f['number_format'] != dominant_format:
                                is_date = 'm' in dominant_format.lower() or 'd' in dominant_format.lower() or 'y' in dominant_format.lower()
                                issue_type = "INCONSISTENT_DATE_FORMAT" if is_date else "INCONSISTENT_NUMBER_FORMAT"
                                
                                # Safety: formulas and merged cells are UNSAFE
                                classification = IssueClassification.AUTO_FIXABLE
                                if f['is_merged'] or f['is_formula']:
                                    classification = IssueClassification.UNSAFE_TO_FIX
                                else:
                                    # Might want review if we aren't completely sure, but AUTO_FIXABLE is fine for format norm.
                                    classification = IssueClassification.AUTO_FIXABLE
                                    
                                issues.append(StandardizedIssue(
                                    issue_id=f"EXCEL_{issue_type}_{col_name}_{f['row_index']}",
                                    source="excel_quality",
                                    issue_type=issue_type,
                                    rule_or_check_name="dominant_pattern",
                                    row_index=f['row_index'],
                                    column=col_name,
                                    actual_value=f['number_format'],
                                    expected_condition=dominant_format,
                                    severity="MINOR",
                                    message=f"Cell has format '{f['number_format']}', expected dominant format '{dominant_format}'.",
                                    recommendation=f"normalize_number_format",
                                    classification=classification
                                ))
                
            # 2. Check Fill
            fills = [f['fill'] for f in populated]
            counter = Counter(fills)
            if len(counter) > 1:
                dominant_fill, count = counter.most_common(1)[0]
                if count / len(populated) >= 0.90:
                    for f in populated:
                        if f['fill'] != dominant_fill:
                            # Fill differences are usually REVIEW_REQUIRED because they could be intentional highlights
                            classification = IssueClassification.REVIEW_REQUIRED
                            if f['is_merged'] or f['is_formula']:
                                classification = IssueClassification.UNSAFE_TO_FIX
                                
                            issues.append(StandardizedIssue(
                                issue_id=f"EXCEL_INCONSISTENT_FILL_{col_name}_{f['row_index']}",
                                source="excel_quality",
                                issue_type="INCONSISTENT_FILL",
                                rule_or_check_name="dominant_pattern",
                                row_index=f['row_index'],
                                column=col_name,
                                actual_value=str(f['fill']),
                                expected_condition=str(dominant_fill),
                                severity="MINOR",
                                message="Cell fill color differs from dominant column pattern.",
                                recommendation="normalize_fill",
                                classification=classification
                            ))
                            
            # 3. Check Font (Bold)
            bolds = [f['font_bold'] for f in populated]
            counter = Counter(bolds)
            if len(counter) > 1:
                dominant_bold, count = counter.most_common(1)[0]
                if count / len(populated) >= 0.95:
                    for f in populated:
                        if f['font_bold'] != dominant_bold:
                            classification = IssueClassification.REVIEW_REQUIRED
                            if f['is_merged'] or f['is_formula']:
                                classification = IssueClassification.UNSAFE_TO_FIX
                                
                            issues.append(StandardizedIssue(
                                issue_id=f"EXCEL_INCONSISTENT_FONT_{col_name}_{f['row_index']}",
                                source="excel_quality",
                                issue_type="INCONSISTENT_FONT",
                                rule_or_check_name="dominant_pattern",
                                row_index=f['row_index'],
                                column=col_name,
                                actual_value=f"Bold: {f['font_bold']}",
                                expected_condition=f"Bold: {dominant_bold}",
                                severity="MINOR",
                                message="Cell font weight differs from dominant column pattern.",
                                recommendation="normalize_font",
                                classification=classification
                            ))
                            
            # 4. Check Alignment
            aligns = [f['alignment_horz'] for f in populated]
            counter = Counter(aligns)
            if len(counter) > 1:
                dominant_align, count = counter.most_common(1)[0]
                if count / len(populated) >= 0.95:
                    for f in populated:
                        if f['alignment_horz'] != dominant_align:
                            classification = IssueClassification.AUTO_FIXABLE
                            if f['is_merged'] or f['is_formula']:
                                classification = IssueClassification.UNSAFE_TO_FIX
                                
                            issues.append(StandardizedIssue(
                                issue_id=f"EXCEL_INCONSISTENT_ALIGNMENT_{col_name}_{f['row_index']}",
                                source="excel_quality",
                                issue_type="INCONSISTENT_ALIGNMENT",
                                rule_or_check_name="dominant_pattern",
                                row_index=f['row_index'],
                                column=col_name,
                                actual_value=str(f['alignment_horz']),
                                expected_condition=str(dominant_align),
                                severity="MINOR",
                                message="Cell horizontal alignment differs from dominant column pattern.",
                                recommendation="normalize_alignment",
                                classification=classification
                            ))
                            
        wb.close()
        return issues
