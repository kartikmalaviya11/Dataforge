import io
import pandas as pd
from typing import List, Dict, Any, Optional
from datetime import datetime

from .repair_engine import RepairResult, ChangeRecord
from .proposal_generator import RepairProposal
from .revalidation import ComparisonResult
from .issue_manager import StandardizedIssue

class ExportService:
    """
    Export Layer (Phase 5).
    Responsible for formatting and serializing the repaired working dataset 
    along with audit/review/summary sheets into CSV and Excel formats.
    """
    
    def export_csv(self, df: pd.DataFrame) -> bytes:
        """
        Exports the cleaned DataFrame to CSV format.
        Preserves column order and handles null values consistently.
        """
        return df.to_csv(index=False, encoding='utf-8-sig').encode('utf-8')
        
    def export_excel(self, 
                     df: pd.DataFrame, 
                     repair_results: Dict[str, RepairResult],
                     proposals: List[RepairProposal],
                     comparison_result: ComparisonResult,
                     after_issues: List[StandardizedIssue],
                     raw_bytes: Optional[bytes] = None,
                     filename: Optional[str] = None) -> bytes:
        """
        Exports the cleaned DataFrame and audit sheets to an Excel workbook.
        If original workbook bytes are provided (Phase 6), it preserves the workbook structure.
        """
        output = io.BytesIO()
        
        # Build audit dataframes
        changes_data = []
        if repair_results:
            for req_action, res in repair_results.items():
                if res.success:
                    for record in res.change_records:
                        changes_data.append({
                            'Row': record.row_index,
                            'Column': record.column,
                            'Old Value': record.old_value,
                            'New Value': record.new_value,
                            'Action': record.action,
                            'Reason': record.reason,
                            'Classification': record.classification
                        })
        
        changes_df = pd.DataFrame(changes_data)
        if changes_df.empty:
            changes_df = pd.DataFrame(columns=['Row', 'Column', 'Old Value', 'New Value', 'Action', 'Reason', 'Classification'])
            
        review_data = []
        for p in proposals:
            if p.approval_state.value != "APPROVED":
                review_data.append({
                    'Row': p.row_index,
                    'Column': p.column,
                    'Issue': p.reason,
                    'Classification': p.classification.value,
                    'Suggested Action': p.action,
                    'Status': p.approval_state.value
                })
                
        for issue in after_issues:
            if issue.classification.value == "UNSAFE_TO_FIX":
                review_data.append({
                    'Row': issue.row_index,
                    'Column': issue.column,
                    'Issue': issue.message,
                    'Classification': issue.classification.value,
                    'Suggested Action': issue.recommendation,
                    'Status': 'UNSAFE'
                })
                
        review_df = pd.DataFrame(review_data)
        if review_df.empty:
            review_df = pd.DataFrame(columns=['Row', 'Column', 'Issue', 'Classification', 'Suggested Action', 'Status'])
            
        summary_df = pd.DataFrame()
        if comparison_result:
            repairs_applied = sum(1 for p in proposals if p.approval_state.value == "APPROVED")
            repairs_rejected = sum(1 for p in proposals if p.approval_state.value == "REJECTED")
            unsafe_issues = sum(1 for iss in after_issues if iss.classification.value == "UNSAFE_TO_FIX")
            
            summary_data = {
                'Metric': [
                    'Issues Before', 'Issues After', 'Issues Resolved', 'Issues Remaining', 'New Issues',
                    'Readiness Before', 'Readiness After', 'Readiness Change',
                    'Repairs Applied', 'Repairs Rejected', 'Unsafe Issues'
                ],
                'Value': [
                    comparison_result.before_issue_count, comparison_result.after_issue_count,
                    comparison_result.resolved_issue_count, comparison_result.remaining_issue_count,
                    comparison_result.new_issue_count,
                    f"{comparison_result.before_readiness:.1f}%", f"{comparison_result.after_readiness:.1f}%",
                    f"{comparison_result.readiness_change:+.1f}%",
                    repairs_applied, repairs_rejected, unsafe_issues
                ]
            }
            summary_df = pd.DataFrame(summary_data)
            
        is_excel_native = raw_bytes and filename and filename.lower().endswith(('.xlsx', '.xls'))
        
        if is_excel_native:
            import openpyxl
            from openpyxl.utils.dataframe import dataframe_to_rows
            from openpyxl.styles import PatternFill, Font, Alignment
            
            # Phase 6: Load original workbook to preserve formulas, formatting, sheets, merged cells
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
            ws = wb.active # The data sheet
            
            # Build column index mapping
            header_row = [cell.value for cell in ws[1]]
            col_name_to_idx = {name: idx + 1 for idx, name in enumerate(header_row) if name}
            
            # Collect all change records
            all_changes = []
            if repair_results:
                for res in repair_results.values():
                    if res.success:
                        all_changes.extend(res.change_records)
                        
            # Separate deletions from modifications
            deletions = []
            modifications = []
            for cr in all_changes:
                if cr.action == "remove_exact_duplicates":
                    deletions.append(cr)
                else:
                    modifications.append(cr)
                    
            # Apply modifications first
            for cr in modifications:
                if cr.row_index is not None and cr.column in col_name_to_idx:
                    excel_row = cr.row_index + 2
                    col_idx = col_name_to_idx[cr.column]
                    cell = ws.cell(row=excel_row, column=col_idx)
                    
                    if cr.action == "normalize_number_format":
                        cell.number_format = cr.new_value
                    elif cr.action == "normalize_date_display_format":
                        cell.number_format = cr.new_value
                    elif cr.action == "normalize_fill":
                        cell.fill = PatternFill(start_color=cr.new_value, end_color=cr.new_value, fill_type="solid")
                    elif cr.action == "normalize_font":
                        # Assume new_value is a dict or string we parse, for now just handle bold
                        if isinstance(cr.new_value, bool) or str(cr.new_value).lower() in ["true", "false"]:
                            bold_val = cr.new_value is True or str(cr.new_value).lower() == "true"
                            current_font = cell.font
                            if current_font:
                                cell.font = Font(name=current_font.name, size=current_font.size, bold=bold_val, italic=current_font.italic)
                            else:
                                cell.font = Font(bold=bold_val)
                    elif cr.action == "normalize_alignment":
                        current_align = cell.alignment
                        if current_align:
                            cell.alignment = Alignment(horizontal=cr.new_value, vertical=current_align.vertical)
                        else:
                            cell.alignment = Alignment(horizontal=cr.new_value)
                    else:
                        # Standard value repair (e.g. trim_whitespace)
                        cell.value = cr.new_value
                        
            # Apply deletions bottom-up to prevent index shifting
            if deletions:
                # Sort descending
                deletions.sort(key=lambda x: x.row_index, reverse=True)
                for cr in deletions:
                    excel_row = cr.row_index + 2
                    ws.delete_rows(excel_row)
                    
            # Append Audit sheets natively
            def add_sheet_from_df(wb, title, df):
                sheet = wb.create_sheet(title=title)
                for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True)):
                    for c_idx, value in enumerate(row, 1):
                        sheet.cell(row=r_idx + 1, column=c_idx, value=value)
            
            add_sheet_from_df(wb, "Changes_Made", changes_df)
            add_sheet_from_df(wb, "Needs_Review", review_df)
            if not summary_df.empty:
                add_sheet_from_df(wb, "Summary", summary_df)
                
            wb.save(output)
            
        else:
            # Fallback to pure pandas export (Phase 5 default)
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Cleaned_Data', index=False)
                changes_df.to_excel(writer, sheet_name='Changes_Made', index=False)
                review_df.to_excel(writer, sheet_name='Needs_Review', index=False)
                if not summary_df.empty:
                    summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
        return output.getvalue()
