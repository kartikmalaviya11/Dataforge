import pytest
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
import io

from src.excel_quality import ExcelQualityEngine
from src.issue_manager import IssueClassification
from src.repair_engine import RepairEngine, RepairInstruction

@pytest.fixture
def sample_excel_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    
    # Headers
    ws.append(["Sales", "Date", "Status", "Merged", "Formula"])
    
    # Rows 2 to 21 (20 data rows)
    for i in range(2, 22):
        # 1. Number formats
        sales_cell = ws.cell(row=i, column=1, value=100.5)
        if i == 5:
            sales_cell.number_format = "General"
        else:
            sales_cell.number_format = "$#,##0.00"
            
        # 2. Date formats
        date_cell = ws.cell(row=i, column=2, value="2023-01-01")
        if i == 6:
            date_cell.number_format = "mm/dd/yyyy"
        else:
            date_cell.number_format = "yyyy-mm-dd"
            
        # 3. Fill
        status_cell = ws.cell(row=i, column=3, value="OK")
        if i == 7:
            status_cell.fill = PatternFill(start_color="FF0000", fill_type="solid")
        else:
            status_cell.fill = PatternFill(start_color="00FF00", fill_type="solid")
            
        # 4. Merged
        ws.cell(row=i, column=4, value="M")
        
        # 5. Formula
        form_cell = ws.cell(row=i, column=5, value=f"=A{i}*2")
        # Let's say formula cell has a format anomaly
        if i == 8:
            form_cell.number_format = "General"
        else:
            form_cell.number_format = "$#,##0.00"
            
    # Merge cells 4:2 to 4:3
    ws.merge_cells(start_row=2, start_column=4, end_row=3, end_column=4)
    # The merged cell at row 2 has a fill anomaly
    ws.cell(row=2, column=4).fill = PatternFill(start_color="FF0000", fill_type="solid")
    for i in range(4, 22):
        ws.cell(row=i, column=4).fill = PatternFill(start_color="00FF00", fill_type="solid")
    
    wb.create_sheet("EmptySheet")
    
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def test_detection_and_classification(sample_excel_bytes):
    # Tests A, B, C, D, E, F, G, H, I
    df = pd.DataFrame({
        "Sales": [100.5] * 20,
        "Date": ["2023-01-01"] * 20,
        "Status": ["OK"] * 20,
        "Merged": ["M"] * 20,
        "Formula": [201.0] * 20
    })
    
    engine = ExcelQualityEngine(sample_excel_bytes)
    issues = engine.analyze(df)
    
    issue_types = {i.issue_type for i in issues}
    assert "INCONSISTENT_NUMBER_FORMAT" in issue_types
    assert "INCONSISTENT_DATE_FORMAT" in issue_types
    assert "INCONSISTENT_FILL" in issue_types
    
    # Check classifications
    num_issue = next(i for i in issues if i.issue_type == "INCONSISTENT_NUMBER_FORMAT" and i.column == "Sales")
    assert num_issue.classification == IssueClassification.AUTO_FIXABLE # G: Clear deterministic mismatch
    
    date_issue = next(i for i in issues if i.issue_type == "INCONSISTENT_DATE_FORMAT")
    assert date_issue.classification == IssueClassification.AUTO_FIXABLE
    
    fill_issue = next(i for i in issues if i.issue_type == "INCONSISTENT_FILL" and i.column == "Status")
    assert fill_issue.classification == IssueClassification.REVIEW_REQUIRED # H: Likely but review-worthy
    
    formula_issue = next(i for i in issues if i.issue_type == "INCONSISTENT_NUMBER_FORMAT" and i.column == "Formula")
    assert formula_issue.classification == IssueClassification.UNSAFE_TO_FIX # I: Formula cells are unsafe
    
    merged_issue = next((i for i in issues if i.issue_type == "INCONSISTENT_FILL" and i.column == "Merged"), None)
    if merged_issue:
        assert merged_issue.classification == IssueClassification.UNSAFE_TO_FIX # I: Merged cells are unsafe

def test_repair_operations(sample_excel_bytes):
    # Tests J, K, L, M, N, O, P, Q, R, S, T, U, V
    df = pd.DataFrame({
        "Sales": [100.5] * 20,
        "Date": ["2023-01-01"] * 20,
        "Status": ["OK"] * 20,
        "Merged": ["M"] * 20,
        "Formula": [201.0] * 20
    })
    
    # Instruction for Sales format (AUTO_FIXABLE -> APPROVED)
    inst1 = RepairInstruction(
        action="normalize_number_format",
        classification=IssueClassification.AUTO_FIXABLE,
        column="Sales",
        row_index=3, # excel_row = 5
        proposed_value="$#,##0.00",
        approved=True
    )
    
    repair_engine = RepairEngine()
    repaired_df, res_dict = repair_engine.apply(df, [inst1])
    
    # df remains identical (values unchanged)
    pd.testing.assert_frame_equal(df, repaired_df)
    
    assert res_dict["instruction_0"].success
    
    # Let's run ExportService to apply the changes to the workbook!
    from src.export_service import ExportService
    from src.revalidation import ComparisonResult
    
    exporter = ExportService()
    comp = ComparisonResult(0,0,0,0,0,100,100,0,[],[],[])
    excel_bytes = exporter.export_excel(
        repaired_df, res_dict, [], comp, [], raw_bytes=sample_excel_bytes, filename="test.xlsx"
    )
    
    # Check the resulting workbook integrity and repairs
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    
    # O, P: Sheets preserved
    assert len(wb.sheetnames) >= 2
    assert "Data" in wb.sheetnames
    assert "EmptySheet" in wb.sheetnames
    assert "Changes_Made" in wb.sheetnames
    
    ws = wb["Data"]
    
    # J: Number format normalized
    assert ws.cell(row=5, column=1).number_format == "$#,##0.00"
    
    # R: Formula remains formula
    assert ws.cell(row=8, column=5).data_type == 'f'
    assert str(ws.cell(row=8, column=5).value).startswith('=A8')
    
    # Q: Merged ranges preserved
    assert len(ws.merged_cells.ranges) > 0
    
    # S: Data values remain unchanged
    assert ws.cell(row=5, column=1).value == 100.5
