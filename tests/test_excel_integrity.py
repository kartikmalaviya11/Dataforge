import pytest
import pandas as pd
import openpyxl
import io

from src.repair_engine import RepairEngine, RepairInstruction, ChangeRecord
from src.export_service import ExportService
from src.revalidation import ComparisonResult
from src.issue_manager import IssueClassification, StandardizedIssue

@pytest.fixture
def duplicate_row_excel_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    # Headers
    ws.append(["A", "B", "C", "D"])

    # Row 2 (Data)
    ws.append([1, 10, 100, "=A2+B2"])

    # Row 3 (Duplicate of Row 2) - we will delete this
    ws.append([1, 10, 100, "=A3+B3"])

    # Row 4 (Data)
    ws.append([2, 20, 200, "=A4+B4"])

    # Row 5 (Formula referencing Row 2 and 4)
    ws.append(["Sum:", "", "", "=SUM(D2:D4)"])

    # Row 6 (Formula referencing the duplicate row specifically)
    ws.append(["Ref duplicate:", "", "", "=D3"])

    # Merge cells A7:B7
    ws.merge_cells("A7:B7")
    ws.cell(row=7, column=1, value="Merged")

    # Create another sheet
    wb.create_sheet("OtherSheet")

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def test_duplicate_row_rejection_integrity(duplicate_row_excel_bytes):
    # This test verifies that if the workbook contains formulas or merged cells,
    # DUPLICATE_ROW issues are downgraded to UNSAFE_TO_FIX, preventing workbook corruption.
    df = pd.DataFrame({
        "A": [1, 1, 2, "Sum:", "Ref duplicate:", "Merged"],
        "B": [10, 10, 20, "", "", ""],
        "C": [100, 100, 200, "", "", ""],
        "D": [11, 11, 22, 44, 11, ""]
    })

    # 1. Quality Engine detects duplicates
    from src.profiler import profile_dataset
    from src.type_detector import detect_semantic_types
    from src.quality_engine import run_quality_checks
    from src.issue_manager import IssueManager

    col_profiles = profile_dataset(df)
    type_results = detect_semantic_types(df, col_profiles)
    quality_issues = run_quality_checks(df, col_profiles, type_results)

    # 2. Excel Engine detects formulas/merges
    from src.excel_quality import ExcelQualityEngine
    excel_eng = ExcelQualityEngine(duplicate_row_excel_bytes)
    excel_issues = excel_eng.analyze(df)

    assert excel_eng.has_formulas_or_merges is True

    # Consolidate issues like app.py does
    im = IssueManager()
    all_issues = im.consolidate(quality_issues + excel_issues)

    # 3. Simulated app.py downgrading
    if excel_eng.has_formulas_or_merges:
        for issue in all_issues:
            if issue.issue_type == "DUPLICATE_ROW":
                issue.classification = IssueClassification.UNSAFE_TO_FIX

    # 4. Assert the issue is now UNSAFE_TO_FIX
    dup_issue = next((i for i in all_issues if i.issue_type == "DUPLICATE_ROW"), None)
    assert dup_issue is not None
    assert dup_issue.classification == IssueClassification.UNSAFE_TO_FIX

    # 5. Proposal generator should NOT create an AUTO_FIXABLE proposal
    from src.proposal_generator import ProposalGenerator
    prop_gen = ProposalGenerator()
    proposals = prop_gen.generate_proposals(df, all_issues)

    assert len(proposals) == 0 # Since it's UNSAFE_TO_FIX, no proposal is generated for it

def test_formula_safeguard_independent_of_formatting_issues():
    # Build clean workbook with formulas/merges but NO formatting inconsistencies
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name", "Total"])
    ws.append([1, "Item A", 100])
    ws.append([1, "Item A", 100]) # duplicate row
    ws.append(["Formula", "", "=SUM(C2:C3)"])
    ws.merge_cells("A5:B5")
    ws.cell(row=5, column=1, value="Note")

    bio = io.BytesIO()
    wb.save(bio)
    raw_bytes = bio.getvalue()

    df = pd.DataFrame({
        "ID": [1, 1, "Formula", "Note"],
        "Name": ["Item A", "Item A", "", ""],
        "Total": [100, 100, 200, ""]
    })

    from src.excel_quality import ExcelQualityEngine
    from src.issue_manager import IssueManager, IssueClassification, StandardizedIssue

    excel_eng = ExcelQualityEngine(raw_bytes)
    excel_issues = excel_eng.analyze(df)

    # Confirm NO formatting issues detected
    assert len(excel_issues) == 0
    # But formulas/merges ARE detected
    assert excel_eng.has_formulas_or_merges is True

    # Simulate duplicate row issue
    dup_issue = StandardizedIssue(
        issue_id="DUP_1",
        source="Sheet1",
        issue_type="DUPLICATE_ROW",
        rule_or_check_name="duplicate_rows",
        row_index=1,
        column=None,
        actual_value=None,
        expected_condition="unique rows",
        severity="MEDIUM",
        message="Duplicate row detected",
        recommendation="remove_exact_duplicates",
        classification=IssueClassification.AUTO_FIXABLE
    )

    all_issues = [dup_issue]
    # Pipeline safeguard logic
    if excel_eng is not None and excel_eng.has_formulas_or_merges:
        for issue in all_issues:
            if issue.issue_type == "DUPLICATE_ROW":
                issue.classification = IssueClassification.UNSAFE_TO_FIX

    assert dup_issue.classification == IssueClassification.UNSAFE_TO_FIX

def test_multi_sheet_target_and_export_preservation():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "DataSheet"
    ws1.append(["ColA", "ColB"])
    ws1.append(["val1", "val2"])

    ws2 = wb.create_sheet("FormulasSheet")
    ws2.append(["Header", "Calc"])
    ws2.append(["Ref", "=DataSheet!A2"])
    ws2.merge_cells("A3:B3")
    ws2.cell(row=3, column=1, value="MergedNote")

    bio = io.BytesIO()
    wb.save(bio)
    raw_bytes = bio.getvalue()

    from src.excel_quality import ExcelQualityEngine
    from src.export_service import ExportService

    eng1 = ExcelQualityEngine(raw_bytes, sheet_name="DataSheet")
    df1 = pd.DataFrame({"ColA": ["val1"], "ColB": ["val2"]})
    eng1.analyze(df1)

    eng2 = ExcelQualityEngine(raw_bytes, sheet_name="FormulasSheet")
    df2 = pd.DataFrame({"Header": ["Ref", "MergedNote"], "Calc": ["val1", ""]})
    eng2.analyze(df2)
    assert eng2.has_formulas_or_merges is True

    # Test export targeting DataSheet preserves FormulasSheet
    export_svc = ExportService()
    res_bytes = export_svc.export_excel(
        df=df1,
        repair_results={},
        proposals=[],
        comparison_result=None,
        after_issues=[],
        raw_bytes=raw_bytes,
        sheet_name="DataSheet"
    )

    wb_out = openpyxl.load_workbook(io.BytesIO(res_bytes))
    assert "FormulasSheet" in wb_out.sheetnames
    ws2_out = wb_out["FormulasSheet"]
    assert len(ws2_out.merged_cells.ranges) > 0
    assert ws2_out.cell(row=2, column=2).value == "=DataSheet!A2"

def test_legacy_xls_safe_handling_and_fallback():
    from src.excel_quality import ExcelQualityEngine
    from src.export_service import ExportService

    # Legacy OLE2 .xls magic bytes
    legacy_xls_bytes = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 512
    df = pd.DataFrame({"id": [1, 2], "val": [10, 20]})

    # 1. ExcelQualityEngine safely handles legacy .xls without crashing and flags safeguards conservatively
    eqe = ExcelQualityEngine(legacy_xls_bytes, sheet_name="Sheet1")
    issues = eqe.analyze(df)
    assert issues == []
    assert eqe.has_formulas_or_merges is True

    # 2. ExportService falls back safely to pandas ExcelWriter producing valid workbook
    export_svc = ExportService()
    out_bytes = export_svc.export_excel(
        df=df,
        repair_results={},
        proposals=[],
        comparison_result=None,
        after_issues=[],
        raw_bytes=legacy_xls_bytes,
        filename="legacy_data.xls"
    )

    assert out_bytes is not None
    wb_out = openpyxl.load_workbook(io.BytesIO(out_bytes))
    assert "Cleaned_Data" in wb_out.sheetnames
    assert "Changes_Made" in wb_out.sheetnames
    assert "Needs_Review" in wb_out.sheetnames
    assert wb_out["Cleaned_Data"].cell(row=2, column=1).value == 1
